# -*- coding: utf-8 -*-
"""Lecture des fichiers déposés par l'utilisateur et aperçus (Excel / PDF)."""

from __future__ import annotations

import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

MAX_ROWS = 2000            # lignes lues par feuille (au-delà : tronqué, signalé)
MAX_TEXT = 40_000          # caractères de texte transmis au modèle


def safe_filename(name: str, ext: str) -> str:
    base = re.sub(r"[^\w\-. ]+", "_", str(name or "fichier")).strip(" ._") or "fichier"
    base = re.sub(r"\.(xlsx|pdf|csv|docx)$", "", base, flags=re.I)[:60]
    return f"{base}.{ext}"


def read_upload(name: str, data: bytes) -> dict:
    """-> {"name", "kind": "table"|"text", "tables": [{"name","columns","rows"}], "text": str, "truncated": bool}"""
    low = name.lower()
    out = {"name": name, "kind": "text", "tables": [], "text": "", "truncated": False}
    try:
        if low.endswith((".xlsx", ".xlsm", ".xls", ".csv")):
            import pandas as pd
            if low.endswith(".csv"):
                sheets = {"CSV": pd.read_csv(io.BytesIO(data), sep=None, engine="python", dtype=object)}
            else:
                sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=object)
            for sh, df in sheets.items():
                df = df.dropna(how="all").dropna(axis=1, how="all")
                if df.empty:
                    continue
                if len(df) > MAX_ROWS:
                    out["truncated"] = True
                    df = df.head(MAX_ROWS)
                out["tables"].append({"name": str(sh), "columns": [str(c) for c in df.columns],
                                      "rows": [[None if (v is None or v != v) else v for v in r] for r in df.itertuples(index=False)]})
            out["kind"] = "table"
        elif low.endswith(".docx"):
            import docx
            d = docx.Document(io.BytesIO(data))
            out["text"] = "\n".join(p.text for p in d.paragraphs if p.text.strip())
            for t in d.tables:
                rows = [[c.text.strip() for c in r.cells] for r in t.rows]
                if len(rows) > 1:
                    out["tables"].append({"name": "Tableau Word", "columns": rows[0], "rows": rows[1:]})
        elif low.endswith(".pdf"):
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                parts = [(p.extract_text() or "") for p in pdf.pages[:60]]
            out["text"] = "\n\n".join(parts)
        else:
            out["text"] = data.decode("utf-8", errors="ignore")
    except Exception as exc:  # noqa: BLE001
        out["text"] = f"[Fichier illisible : {type(exc).__name__}: {exc}]"
    if len(out["text"]) > MAX_TEXT:
        out["text"], out["truncated"] = out["text"][:MAX_TEXT], True
    return out


def uploads_to_context(files: list[dict]) -> str:
    """Contexte texte transmis au modèle (données exactes : le modèle doit les recopier, pas les inventer)."""
    parts = []
    for f in files:
        parts.append(f"=== Fichier : {f['name']} ===")
        for t in f["tables"]:
            parts.append(f"[Tableau « {t['name']} » — {len(t['rows'])} ligne(s)]")
            parts.append(" | ".join(t["columns"]))
            for r in t["rows"]:
                parts.append(" | ".join("" if v is None else str(v) for v in r))
        if f["text"]:
            parts.append(f["text"])
        if f["truncated"]:
            parts.append("[… tronqué : fichier plus long que la limite lue]")
    return "\n".join(parts)


# ------------------------------------------------------------------ aperçus
def libreoffice_path() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def recalc_xlsx(data: bytes, timeout: int = 90) -> bytes | None:
    """Recalcule le classeur avec LibreOffice (si installé) : renvoie un xlsx avec les valeurs calculées, sinon None."""
    exe = libreoffice_path()
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as d:
        src, out = Path(d) / "in.xlsx", Path(d) / "out"
        src.write_bytes(data)
        out.mkdir()
        try:
            subprocess.run([exe, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                           check=True, timeout=timeout, capture_output=True)
        except Exception:  # noqa: BLE001
            return None
        res = out / "in.xlsx"
        return res.read_bytes() if res.is_file() else None


def pdf_first_page_png(data: bytes, scale: float = 1.4) -> bytes | None:
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(data)
        img = pdf[0].render(scale=scale).to_pil()
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        return buf.getvalue()
    except Exception:  # noqa: BLE001
        return None
