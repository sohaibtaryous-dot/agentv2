# -*- coding: utf-8 -*-
"""Agent Excel : le modèle décrit le classeur (JSON strict) ; le code le construit avec openpyxl (styles, formules, formats).
Aucun code généré n'est exécuté : seul le JSON validé est interprété."""

from __future__ import annotations

import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from utils import llm as L

FORMATS = {"text": "@", "number": "#,##0.00", "integer": "#,##0", "currency": '#,##0.00 "€"', "percent": "0.0%",
           "date": "DD/MM/YYYY", "general": "General"}
NEW_FUNCS = ("XLOOKUP", "XMATCH", "IFS", "MAXIFS", "MINIFS", "CONCAT", "TEXTJOIN", "FILTER", "UNIQUE", "SORT", "SWITCH")

SCHEMA = {
    "type": "object",
    "properties": {
        "filename": {"type": "string"},
        "sheets": {"type": "array", "minItems": 1, "items": {
            "type": "object",
            "properties": {
                "name": {"type": "string", "description": "Nom de feuille (max 31 caractères, sans : \\ / ? * [ ])"},
                "columns": {"type": "array", "minItems": 1, "items": {
                    "type": "object",
                    "properties": {"header": {"type": "string"}, "width": {"type": "number"},
                                   "format": {"type": "string", "enum": list(FORMATS)}},
                    "required": ["header"]}},
                "rows": {"type": "array", "items": {"type": "array", "items": {"type": ["string", "number", "null"]}},
                         "description": "Données ligne par ligne. Formule = chaîne commençant par « = » (syntaxe anglaise, séparateur virgule)."},
                "freeze_header": {"type": "boolean"}, "autofilter": {"type": "boolean"},
                "header_color": {"type": "string", "description": "Couleur hex de l'en-tête, ex. 2E7D32"},
            },
            "required": ["name", "columns", "rows"]}},
        "notes": {"type": "string"},
    },
    "required": ["filename", "sheets"],
}

SYSTEM = """Tu es l'agent Excel d'une application de transit / douane. Tu décris un classeur Excel en JSON (schéma imposé).
RÈGLES :
- Ligne 1 = en-têtes (les colonnes), données à partir de la ligne 2. Les formules utilisent donc ces numéros de ligne
  (ex. première ligne de données = 2, 3 lignes de données => lignes 2 à 4, ligne de total = 5).
- Formules : chaîne qui commence par « = », fonctions anglaises (SUM, AVERAGE, IF, VLOOKUP, XLOOKUP, INDEX, MATCH, SUMIF…), arguments séparés par des virgules.
  Utilise de vraies formules pour les totaux, moyennes, pourcentages, écarts, recherches (jamais la valeur calculée en dur).
- Références croisées : 'Nom feuille'!B2:B20. Les plages doivent couvrir exactement les lignes existantes.
- Recopie EXACTEMENT les données fournies (nombres, références, noms). N'invente aucune donnée ; si une valeur manque, laisse null.
- Choisis le format de chaque colonne (text, number, integer, currency, percent, date). Dates : « JJ/MM/AAAA » en texte ISO AAAA-MM-JJ.
- Ajoute une ligne TOTAL avec formules quand c'est pertinent. Noms de feuilles courts, sans caractères interdits.
- Si l'utilisateur demande une modification, pars du classeur précédent (fourni) et ne change que ce qui est demandé."""


def generate_spec(llm: L.LLM, request: str, context: str, previous: dict | None, errors: list[str] | None, history: str = "") -> dict:
    import json
    parts = []
    if history:
        parts.append(f"Historique de la conversation :\n{history}")
    if context:
        parts.append(f"Données fournies par l'utilisateur :\n{context}")
    if previous:
        parts.append("Classeur précédent (JSON) à modifier :\n" + json.dumps(previous, ensure_ascii=False))
    if errors:
        parts.append("Le contrôle qualité a trouvé ces erreurs dans ta version précédente — corrige-les toutes :\n- " + "\n- ".join(errors))
    parts.append(f"Demande : {request}")
    return llm.structured(SYSTEM, "\n\n".join(parts), SCHEMA, "excel_workbook")


def _clean_sheet_name(name: str, used: set) -> str:
    n = re.sub(r"[:\\/?*\[\]]", "-", str(name or "Feuille")).strip("' ")[:31] or "Feuille"
    base, i = n, 2
    while n.lower() in used:
        n = f"{base[:28]}_{i}"
        i += 1
    used.add(n.lower())
    return n


def _fx(value):
    """Formule -> syntaxe fichier (préfixe _xlfn. pour les fonctions récentes, sinon Excel affiche #NAME?)."""
    if isinstance(value, str) and value.startswith("="):
        for fn in NEW_FUNCS:
            value = re.sub(rf"(?<![\w.]){fn}\(", f"_xlfn.{fn}(", value, flags=re.I)
    return value


def build_xlsx(spec: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    default_head = "1F3864"
    band = PatternFill("solid", start_color="EEF3FB", end_color="EEF3FB")
    thin = Side(style="thin", color="B7C2D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for sh in spec["sheets"]:
        ws = wb.create_sheet(_clean_sheet_name(sh.get("name"), used))
        hc = re.sub(r"[^0-9A-Fa-f]", "", str(sh.get("header_color") or ""))[:6]
        head_fill = PatternFill("solid", start_color=hc if len(hc) == 6 else default_head, end_color=hc if len(hc) == 6 else default_head)
        cols = sh["columns"]
        fmts = [FORMATS.get(c.get("format") or "general", "General") for c in cols]
        for j, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=j, value=str(c.get("header", "")))
            cell.font, cell.fill, cell.border = Font(name="Arial", bold=True, color="FFFFFF", size=10), head_fill, border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 26
        for i, row in enumerate(sh["rows"], 2):
            is_total = bool(row) and isinstance(row[0], str) and row[0].strip().upper().startswith("TOTAL")
            for j in range(1, len(cols) + 1):
                v = _fx(row[j - 1] if j - 1 < len(row) else None)
                if fmts[j - 1] == FORMATS["date"] and isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
                    v = datetime.strptime(v, "%Y-%m-%d").date()
                cell = ws.cell(row=i, column=j, value=v)
                cell.font = Font(name="Arial", size=10, bold=is_total)
                cell.border = border
                if fmts[j - 1] != "General" and not (isinstance(v, str) and not v.startswith("=") and fmts[j - 1] != "@"):
                    cell.number_format = fmts[j - 1]
                if i % 2 == 1 and not is_total:
                    cell.fill = band
                if is_total:
                    cell.fill = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
        for j, c in enumerate(cols, 1):                                   # largeurs : demandée, sinon selon le contenu
            w = c.get("width") or max([len(str(c.get("header", "")))] + [len(str(r[j - 1])) for r in sh["rows"] if j - 1 < len(r) and r[j - 1] is not None and not str(r[j - 1]).startswith("=")] + [8])
            ws.column_dimensions[get_column_letter(j)].width = min(max(float(w) + 2, 9), 60)
        if sh.get("freeze_header", True):
            ws.freeze_panes = "A2"
        if sh.get("autofilter", True) and sh["rows"]:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(sh['rows']) + 1}"
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
