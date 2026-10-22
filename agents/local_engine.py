# -*- coding: utf-8 -*-
"""Moteur LOCAL (sans clé API) : comprend les demandes courantes (français, anglais, darija en lettres latines) et construit les
fichiers directement à partir des données déposées. Il produit les mêmes spécifications (JSON) que les agents IA :
elles passent donc par les mêmes constructeurs (openpyxl / ReportLab) et le même contrôle qualité.

Ce que le moteur local sait faire : choisir la table (ou lire les PDF BL / factures avec le moteur du site), total (SUM), TVA x % (HT, TVA, TTC),
écart entre deux colonnes, tri, synthèse par colonne (SUMIFS), supprimer / renommer une colonne, couleur d'en-tête, titre, rapport PDF
(chiffres clés + tableau + texte). Il ne peut PAS inventer de contenu : sans données déposées, il le dit."""

from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime

from openpyxl.utils import get_column_letter

COLORS = {"vert": "2E7D32", "verte": "2E7D32", "green": "2E7D32", "khdar": "2E7D32", "khder": "2E7D32", "bleu": "1F3864", "blue": "1F3864",
          "zr9": "1F3864", "azrak": "1F3864", "zre9": "1F3864", "rouge": "C62828", "red": "C62828", "7mer": "C62828", "hmer": "C62828",
          "orange": "E65100", "noir": "212121", "black": "212121", "kahl": "212121", "gris": "616161", "grey": "616161", "gray": "616161",
          "violet": "6A1B9A", "purple": "6A1B9A", "jaune": "F9A825", "yellow": "F9A825", "sfer": "F9A825", "marron": "5D4037"}
AMOUNT_RE = re.compile(r"montant|amount|prix|price|valeur|value|total|\bht\b|\bttc\b|co[uû]t|cost", re.I)
QTY_RE = re.compile(r"qt[eé]|quantit|qty|colis|packages|nombre|nb\b", re.I)
WORDS_TOTAL = r"total|somme|sum|jmo3|jm3|majmou3|majmo3|addition"


def _norm(s) -> str:
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s.lower()).strip()


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return None if v != v else float(v)
    s = str(v).strip().replace(" ", "").replace(" ", "")
    if not s or not re.fullmatch(r"[-+]?\d[\d.,]*", s):
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".") if len(s.split(",")[-1]) <= 2 else s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def _cell(v):
    if v is None or (isinstance(v, float) and v != v):
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if hasattr(v, "item"):
        v = v.item()
    n = _num(v)
    if isinstance(v, str) and n is not None and not re.match(r"^0\d", v.strip()) and len(v.strip()) < 15:
        return int(n) if n.is_integer() and "." not in v and "," not in v else n
    return v


def _col_format(header: str, values: list) -> str:
    nums = [_num(v) for v in values if v not in (None, "")]
    if not nums or sum(n is not None for n in nums) < 0.8 * len(nums):
        if values and all(isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v) for v in values if v):
            return "date"
        return "text" if any(isinstance(v, str) for v in values) else "general"
    if all(n is not None and float(n).is_integer() for n in nums if n is not None) and not AMOUNT_RE.search(header):
        return "integer" if QTY_RE.search(header) or len(nums) else "number"
    return "number"


# ----------------------------------------------------------------------------- données
def _tables_from(ex) -> list[dict]:
    out = []
    for f in ex.files:
        for t in f["tables"]:
            out.append({"name": f"{f['name']} · {t['name']}" if len(ex.files) > 1 or len(f["tables"]) > 1 else t["name"],
                        "columns": [str(c) for c in t["columns"]], "rows": [[_cell(v) for v in r] for r in t["rows"]]})
    return out


def _pick_table(request: str, tables: list[dict]) -> dict | None:
    if not tables:
        return None
    r = _norm(request)
    for t in tables:
        if _norm(t["name"]) and _norm(t["name"]).split("·")[-1].strip() in r and len(_norm(t["name"])) > 3:
            return t
    if re.search(r"factur", r):
        for t in tables:
            if "factur" in _norm(t["name"]):
                return t
    if re.search(r"\bbl\b|reference|packing", r):
        for t in tables:
            if re.search(r"\bbl\b|reference", _norm(t["name"])):
                return t
    return max(tables, key=lambda t: len(t["rows"]))


def _find_col(name: str, cols: list[str]) -> int | None:
    n = _norm(name).strip("\"'« »")
    if not n:
        return None
    for i, c in enumerate(cols):
        if _norm(c) == n:
            return i
    for i, c in enumerate(cols):
        if n in _norm(c) or (len(_norm(c)) > 2 and _norm(c) in n):
            return i
    return None


def _quoted(request: str) -> list[str]:
    return [a or b or c for a, b, c in re.findall(r"«\s*([^»]+?)\s*»|\"([^\"]+)\"|'([^']{2,})'", request)]


# ----------------------------------------------------------------------------- commandes
def parse_commands(request: str, cols: list[str]) -> dict:
    r = _norm(request)
    cmd: dict = {}
    if re.search(r"sans total", r):
        cmd["total"] = False
    elif re.search(WORDS_TOTAL, r):
        cmd["total"] = True
    m = re.search(r"tva\s*(?:de|a|=|:)?\s*(\d+(?:[.,]\d+)?)\s*%?", r)
    if m:
        cmd["tva"] = float(m.group(1).replace(",", ".")) / 100
    if re.search(r"ecart|difference|diff\b|farq", r):
        cmd["ecart"] = True
    for kw, hexv in COLORS.items():
        if re.search(rf"\b{kw}\b", r) and re.search(r"en-?tete|entete|header|titre colonne|couleur|colou?r|lawn", r):
            cmd["color"] = hexv
            break
    m = re.search(r"(?:trie[rz]?|trier|sort|rattab|rtb)\s+(?:par|by|b)\s+(.+?)(?:\s+(?:et|puis|avec|en|and)\b|[.,;]|$)", r)
    if m:
        i = _find_col(m.group(1), cols)
        if i is not None:
            cmd["sort"] = (i, bool(re.search(r"decroiss|desc|z.?a|ntl3|plus grand", r)))
    mm = re.search(r"(?:recap\w*|synthese|resume|regroup\w*|group\w*)\s+(?:les\s+)?(?:\w+\s+)?(?:par|by|per)\s+([\w' ]{2,30}?)(?:\s+(?:et|puis|avec|en)\b|[.,;]|$)", r)
    if not mm and re.search(r"recap|synthese|regroup|group", r):
        for cand in re.finditer(r"(?<!tri )(?<!trie )(?<!triee )(?<!trier )(?<!sort )(?:par|by|per)\s+([\w' ]{2,30}?)(?:\s+(?:et|puis|avec|en)\b|[.,;]|$)", r):
            if _find_col(cand.group(1), cols) is not None and not re.search(r"(?:tri\w*|sort)\s+(?:par|by)\s+" + re.escape(cand.group(1)), r):
                mm = cand
                break
    if mm:
        i = _find_col(mm.group(1), cols)
        if i is not None:
            cmd["group"] = i
    for mm in re.finditer(r"(?:supprim\w*|enleve\w*|retire\w*|delete|remove|7yyd|hyd)\s+(?:la\s+)?(?:colonne\s+)?([\w'«» \"]+?)(?:\s+(?:et|puis)\b|[.,;]|$)", r):
        i = _find_col(mm.group(1), cols)
        if i is not None:
            cmd.setdefault("drop", []).append(i)
    for mm in re.finditer(r"(?:renomm\w*|rename)\s+(?:la\s+colonne\s+)?(.+?)\s+(?:en|to|b)\s+(.+?)(?:\s+(?:et|puis)\b|[.,;]|$)", request, re.I):
        i = _find_col(mm.group(1), cols)
        if i is not None:
            cmd.setdefault("rename", []).append((i, mm.group(2).strip(" \"'«»")))
    q = _quoted(request)
    if q:
        cmd["quoted"] = q
    m = re.search(r"(?<![A-Za-z])(?:titre|title|nomme|appelle|smiyto|smito)(?![a-z])\s*:?\s*[«\"']?(.+?)[»\"']?(?:[.,;]|$)", request, re.I)
    if m and len(m.group(1)) < 80:
        cmd["title"] = m.group(1).strip()
    return cmd


def _amount_col(cols: list[str], rows: list[list]) -> int | None:
    best = None
    for i, c in enumerate(cols):
        nums = [_num(r[i]) for r in rows if i < len(r) and r[i] not in (None, "")]
        if nums and sum(n is not None for n in nums) >= 0.8 * len(nums) and AMOUNT_RE.search(c):
            if best is None or re.search(r"ht|montant|amount", c, re.I):
                best = i
    return best


def _numeric_cols(cols: list[str], rows: list[list]) -> list[int]:
    out = []
    for i, c in enumerate(cols):
        nums = [_num(r[i]) for r in rows if i < len(r) and r[i] not in (None, "")]
        if nums and sum(n is not None for n in nums) >= 0.8 * len(nums) and not re.search(r"n[°o]|num|date|code|ref|poste|item|ligne|^id$", _norm(c)):
            out.append(i)
    return out


def _apply(state: dict) -> dict:
    """Construit le classeur depuis l'état (données de base + commandes cumulées) : les formules sont toujours recalculées."""
    cmd = state["cmd"]
    drop = set(cmd.get("drop", []))
    ren = cmd.get("rename", {})
    keep = [i for i, c in enumerate(state["cols"]) if c["header"] not in drop]
    cols = [dict(state["cols"][i], header=ren.get(state["cols"][i]["header"], state["cols"][i]["header"])) for i in keep]
    orig = [state["cols"][i]["header"] for i in keep]
    rows = [[(r[i] if i < len(r) else None) for i in keep] for r in state["rows"]]
    heads = [c["header"] for c in cols]
    if cmd.get("sort") and cmd["sort"][0] in orig:
        i, desc = orig.index(cmd["sort"][0]), cmd["sort"][1]
        rows = sorted(rows, key=lambda r: (r[i] is None, _num(r[i]) if _num(r[i]) is not None else _norm(r[i])), reverse=desc)
    n0, notes = len(cols), []
    first, last = 2, len(rows) + 1
    num_cols = _numeric_cols(heads, rows)
    amt = _amount_col(heads, rows)
    if cmd.get("tva") is not None:
        if amt is None:
            notes.append("TVA ignorée : aucune colonne de montant reconnue.")
        else:
            L = get_column_letter(amt + 1)
            cols += [{"header": f"TVA {cmd['tva'] * 100:g} %", "format": "number"}, {"header": "Montant TTC", "format": "number"}]
            tcol = get_column_letter(n0 + 1)
            for k, r in enumerate(rows):
                r += [f"=ROUND({L}{k + 2}*{cmd['tva']},2)", f"={L}{k + 2}+{tcol}{k + 2}"]
    if cmd.get("ecart"):
        pair = None
        e = cmd["ecart"]
        if isinstance(e, (list, tuple)) and e[0] in orig and e[1] in orig:
            pair = (orig.index(e[0]), orig.index(e[1]))
        elif len(num_cols) >= 2:
            pair = (num_cols[-2], num_cols[-1])
        if pair:
            a_, b_ = get_column_letter(pair[0] + 1), get_column_letter(pair[1] + 1)
            cols.append({"header": "Écart", "format": "number"})
            for k, r in enumerate(rows):
                r.append(f"={a_}{k + 2}-{b_}{k + 2}")
        else:
            notes.append("Écart ignoré : il faut deux colonnes numériques.")
    n_rows = len(rows)
    if cmd.get("total", bool(num_cols)) and rows:
        tot = ["TOTAL"] + [None] * (len(cols) - 1)
        for j in range(1, len(cols)):
            if j in num_cols or j >= n0:
                L = get_column_letter(j + 1)
                tot[j] = f"=SUM({L}{first}:{L}{last})"
        rows = rows + [tot]
    main = {"name": state["name"], "columns": cols, "rows": rows, "freeze_header": True, "autofilter": True}
    if cmd.get("color"):
        main["header_color"] = cmd["color"]
    sheets = [main]
    g = cmd.get("group")
    if g and g in orig and n_rows:
        gi = orig.index(g)
        keys = []
        for r in rows[:n_rows]:
            if r[gi] not in (None, "") and r[gi] not in keys:
                keys.append(r[gi])
        nums = [j for j in num_cols if j != gi]
        if keys and nums and len(keys) <= 500:
            G, sn = get_column_letter(gi + 1), main["name"].replace("'", "''")
            scols = [{"header": heads[gi], "format": "text"}] + [{"header": heads[j], "format": cols[j].get("format", "number")} for j in nums]
            srows = [[k] + [f"=SUMIFS('{sn}'!{get_column_letter(j + 1)}${first}:{get_column_letter(j + 1)}${last},'{sn}'!{G}${first}:{G}${last},$A{i + 2})" for j in nums]
                     for i, k in enumerate(keys)]
            srows.append(["TOTAL"] + [f"=SUM({get_column_letter(c + 2)}2:{get_column_letter(c + 2)}{len(keys) + 1})" for c in range(len(nums))])
            sheets.append({"name": f"Synthèse par {heads[gi]}"[:31], "columns": scols, "rows": srows, "freeze_header": True, "autofilter": False})
        else:
            notes.append("Synthèse ignorée : colonne de regroupement ou colonnes numériques introuvables.")
    return {"filename": cmd.get("title") or state["title"], "sheets": sheets, "notes": " ".join(notes), "_state": state}


def excel_spec(request: str, ex, previous: dict | None = None) -> dict:
    if previous and previous.get("_state") and not ex.files:                      # modification : mêmes données, commandes cumulées
        state = {**previous["_state"], "cmd": dict(previous["_state"]["cmd"])}
    else:
        t = _pick_table(request, _tables_from(ex))
        if t is None:
            raise ValueError("Aucune donnée tabulaire : déposez un fichier Excel / CSV (ou des PDF de BL / factures) pour que le moteur local construise le tableau.")
        fm = [_col_format(c, [r[i] for r in t["rows"] if i < len(r)]) for i, c in enumerate(t["columns"])]
        state = {"cols": [{"header": c, "format": f} for c, f in zip(t["columns"], fm)],
                 "rows": [list(r) + [None] * (len(t["columns"]) - len(r)) for r in t["rows"]],
                 "name": re.sub(r"^.*· ", "", t["name"])[:31] or "Données", "title": re.sub(r"\.\w+$", "", t["name"].split(" · ")[0]), "cmd": {}}
    cmd = state["cmd"]
    ren = cmd.get("rename", {})
    drop = set(cmd.get("drop", []))
    vis = [c["header"] for c in state["cols"] if c["header"] not in drop]
    vis_names = [ren.get(h, h) for h in vis]                                       # noms visibles -> noms d'origine
    new = parse_commands(request, vis_names)
    for k in ("total", "tva", "color", "title"):
        if k in new:
            cmd[k] = new[k]
    if "sort" in new:
        cmd["sort"] = (vis[new["sort"][0]], new["sort"][1])
    if "group" in new:
        cmd["group"] = vis[new["group"]]
    if new.get("drop"):
        cmd["drop"] = sorted(set(cmd.get("drop", [])) | {vis[i] for i in new["drop"]})
    if new.get("rename"):
        cmd["rename"] = {**ren, **{vis[i]: nm for i, nm in new["rename"]}}
    if new.get("ecart"):
        q = [i for i in (_find_col(x, vis_names) for x in new.get("quoted", [])) if i is not None]
        cmd["ecart"] = (vis[q[0]], vis[q[1]]) if len(q) >= 2 else True
    return _apply(state)


def pdf_spec(request: str, ex, previous: dict | None = None) -> dict:
    cmd = parse_commands(request, [])
    tables = _tables_from(ex)
    t = _pick_table(request, tables)
    title = cmd.get("title") or (re.sub(r"\.\w+$", "", ex.files[0]["name"]) if ex.files else "Rapport")
    sections = []
    if t is None:
        text = "\n".join(f["text"] for f in ex.files if f["text"]).strip()
        if not text and previous:
            return {**previous, "title": cmd.get("title") or previous["title"]}
        if not text:
            raise ValueError("Aucune donnée : déposez un fichier (Excel, CSV, PDF, Word, texte) pour que le moteur local construise le PDF.")
        paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()][:40]
        sections.append({"heading": "Contenu", "paragraphs": [re.sub(r"\s+", " ", p)[:1500] for p in paras]})
    else:
        cols, rows = t["columns"], t["rows"]
        nums = _numeric_cols(cols, rows)
        key = [f"{len(rows)} ligne(s) dans « {t['name']} »"]
        for j in nums[:6]:
            vals = [_num(r[j]) for r in rows if j < len(r) and _num(r[j]) is not None]
            if vals:
                key.append(f"Total {cols[j]} : {sum(vals):,.2f}".replace(",", " ").replace(".", ","))
        sections.append({"heading": "Chiffres clés", "bullets": key})
        shown = rows[:300]
        trows = [[("" if v is None else v) for v in r] for r in shown]
        tot = None
        if nums and re.search(WORDS_TOTAL, _norm(request)) or nums:
            tot = ["TOTAL"] + [""] * (len(cols) - 1)
            for j in nums:
                if j == 0:
                    continue
                vals = [_num(r[j]) for r in shown if j < len(r) and _num(r[j]) is not None]
                tot[j] = round(sum(vals), 2)
        if tot and len(rows) <= 300:
            trows.append(tot)
        sections.append({"heading": "Détail", "table": {"columns": cols, "rows": trows, "total_row": bool(tot and len(rows) <= 300)}})
        if len(rows) > 300:
            sections[0]["bullets"].append(f"Le tableau affiche les 300 premières lignes sur {len(rows)}.")
    return {"filename": title, "title": title[:80], "subtitle": datetime.now().strftime("Généré le %d/%m/%Y"), "sections": sections,
            "footer": "TIMAR Tanger — Assistant fichiers"}
